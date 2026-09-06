from flask import Flask, request, redirect, url_for, render_template_string
import sqlite3
import os
from datetime import datetime

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, "treker.db")

app = Flask(__name__)


def get_db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.execute(
        """CREATE TABLE IF NOT EXISTS operations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            type TEXT NOT NULL,
            amount REAL NOT NULL,
            category TEXT NOT NULL,
            date TEXT NOT NULL,
            comment TEXT
        )"""
    )
    conn.commit()
    conn.close()


CATEGORIES = {
    "income": ["Заказы", "Продажа материалов", "Прочее"],
    "expense": ["Материалы", "Зарплата", "Аренда", "Транспорт", "Оборудование", "Прочее"],
}


def month_options():
    months = []
    conn = get_db()
    rows = conn.execute("SELECT DISTINCT substr(date,1,7) AS m FROM operations ORDER BY m DESC").fetchall()
    conn.close()
    for r in rows:
        months.append(r["m"])
    if not months:
        months.append(datetime.now().strftime("%Y-%m"))
    return months


IDX = """
<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>Финансовый учёт работы цеха</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: 'Segoe UI', Arial, sans-serif; background: #0f172a; color: #e2e8f0; min-height: 100vh; }
  .wrap { max-width: 980px; margin: 0 auto; padding: 28px 20px 60px; }
  header { display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 10px; margin-bottom: 22px; }
  h1 { font-size: 26px; }
  .kpi { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px,1fr)); gap: 14px; margin-bottom: 24px; }
  .kpi .card { background: #1e293b; border-radius: 14px; padding: 16px 18px; }
  .kpi .card .label { color: #94a3b8; font-size: 13px; }
  .kpi .card .value { font-size: 24px; font-weight: 700; margin-top: 4px; }
  .kpi .inc { color: #4ade80; } .kpi .exp { color: #f87171; } .kpi .bal { color: #facc15; }
  .panel { background: #1e293b; border-radius: 14px; padding: 16px 18px; margin-bottom: 24px; }
  .panel h2 { font-size: 17px; margin-bottom: 12px; }
  .fields { display: flex; gap: 10px; flex-wrap: wrap; }
  select, input, button { padding: 8px 12px; border-radius: 9px; border: 1px solid #334155; background: #0f172a; color: #e2e8f0; font-size: 14px; }
  button { cursor: pointer; background: linear-gradient(135deg,#0ea5e9,#2563eb); border: none; font-weight: 600; }
  table { width: 100%; border-collapse: collapse; }
  th, td { padding: 9px 10px; text-align: left; border-bottom: 1px solid #334155; font-size: 14px; }
  th { color: #94a3b8; font-weight: 600; }
  .green { color: #4ade80; } .red { color: #f87171; }
  .del { color: #f87171; text-decoration: none; font-weight: 700; }
  .muted { color: #94a3b8; font-size: 13px; }
  .chart { display: flex; align-items: flex-end; gap: 18px; margin-top: 12px; height: 130px; }
  .bar { width: 46px; text-align: center; font-size: 12px; }
  .bar .h { background: #2563eb; border-radius: 8px 8px 0 0; margin-bottom: 4px; }
  .btn-sm { padding: 5px 10px; font-size: 13px; }
</style>
</head>
<body>
<div class="wrap">
  <header>
    <h1>Финансовый учёт работы цеха</h1>
    <form method="get" style="display:flex;gap:8px;align-items:center">
      <label class="muted">Месяц</label>
      <select name="month" onchange="this.form.submit()">
        {% for m in months %}
        <option value="{{ m }}" {% if m == month %}selected{% endif %}>{{ m }}</option>
        {% endfor %}
      </select>
      <a class="btn-sm" style="background:linear-gradient(135deg,#0ea5e9,#2563eb);border-radius:9px;padding:6px 12px;color:#fff;text-decoration:none" href="/export?month={{ month }}">Выгрузить Excel</a>
    </form>
  </header>

  <div class="kpi">
    <div class="card"><div class="label">Доходы</div><div class="value inc">{{ income|int }} ₽</div></div>
    <div class="card"><div class="label">Расходы</div><div class="value exp">{{ expense|int }} ₽</div></div>
    <div class="card"><div class="label">Баланс</div><div class="value bal">{{ balance|int }} ₽</div></div>
  </div>

  <div class="panel">
    <h2>Новая операция</h2>
    <form method="post" action="/add" class="fields">
      <select name="type" required>
        <option value="income">Доход</option>
        <option value="expense">Расход</option>
      </select>
      <select name="category" required>
        <optgroup label="Доходы">{% for c in cats["income"] %}<option>{{ c }}</option>{% endfor %}</optgroup>
        <optgroup label="Расходы">{% for c in cats["expense"] %}<option>{{ c }}</option>{% endfor %}</optgroup>
      </select>
      <input type="number" name="amount" placeholder="Сумма, ₽" step="0.01" min="0" required>
      <input type="date" name="date" value="{{ today }}" required>
      <input type="text" name="comment" placeholder="Комментарий (необязательно)" style="flex:1;min-width:140px">
      <button type="submit">Добавить</button>
    </form>
  </div>

  <div class="panel">
    <h2>Дашборд (по категориям)</h2>
    <div class="chart">
      {% for b in bars %}
      <div class="bar"><div class="h" style="height:{{ b.pct }}%;background:{{ b.color }}">{{ b.val|int }}</div>{{ b.cat }}</div>
      {% endfor %}
    </div>
  </div>

  <div class="panel">
    <h2>Журнал операций ({{ month }})</h2>
    <table>
      <tr><th>Дата</th><th>Тип</th><th>Категория</th><th>Сумма</th><th>Комментарий</th><th>Действия</th></tr>
      {% for o in ops %}
      <tr>
        <td>{{ o["date"] }}</td>
        <td>{% if o["type"]=="income" %}+{% else %}-{% endif %}</td>
        <td>{{ o["category"] }}</td>
        <td class="{% if o['type']=='income' %}green{% else %}red{% endif %}">{{ o["amount"]|int }} ₽</td>
        <td class="muted">{{ o["comment"] or "" }}</td>
        <td><a class="del" href="#" onclick="if(confirm('Удалить операцию?')) location.href='/delete/{{ o['id'] }}'">✕</a></td>
      </tr>
      {% else %}
      <tr><td colspan="6" class="muted">Пока нет операций за этот месяц.</td></tr>
      {% endfor %}
    </table>
  </div>
</div>
</body>
</html>
"""


@app.route("/")
def index():
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    conn = get_db()
    ops = conn.execute(
        "SELECT * FROM operations WHERE substr(date,1,7)=? ORDER BY date DESC, id DESC", (month,)
    ).fetchall()
    income = conn.execute(
        "SELECT COALESCE(SUM(amount),0) AS s FROM operations WHERE substr(date,1,7)=? AND type='income'", (month,)
    ).fetchone()["s"]
    expense = conn.execute(
        "SELECT COALESCE(SUM(amount),0) AS s FROM operations WHERE substr(date,1,7)=? AND type='expense'", (month,)
    ).fetchone()["s"]
    conn.close()
    balance = income - expense
    exp_cats = {}
    for o in ops:
        if o["type"] == "expense":
            exp_cats[o["category"]] = exp_cats.get(o["category"], 0) + o["amount"]
    top = dict(sorted(exp_cats.items(), key=lambda x: x[1], reverse=True)[:5])
    bars = []
    maxval = max(top.values(), default=1)
    colors = ["#2563eb", "#0ea5e9", "#38bdf8", "#818cf8", "#a78bfa"]
    for i, (cat, val) in enumerate(top.items()):
        bars.append({"cat": cat, "val": val, "pct": max(int(val / maxval * 100), 2), "color": colors[i % len(colors)]})
    return render_template_string(
        IDX,
        ops=ops,
        income=income,
        expense=expense,
        balance=balance,
        cats=CATEGORIES,
        months=month_options(),
        month=month,
        today=datetime.now().strftime("%Y-%m-%d"),
        bars=bars,
    )


@app.route("/add", methods=["POST"])
def add():
    t = request.form["type"]
    amount = float(request.form["amount"])
    cat = request.form["category"]
    date = request.form["date"]
    comment = request.form.get("comment", "")
    conn = get_db()
    conn.execute(
        "INSERT INTO operations (type, amount, category, date, comment) VALUES (?,?,?,?,?)",
        (t, amount, cat, date, comment),
    )
    conn.commit()
    conn.close()
    return redirect(url_for("index", month=date[:7]))


@app.route("/export")
def export():
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    conn = get_db()
    ops = conn.execute(
        "SELECT * FROM operations WHERE substr(date,1,7)=? ORDER BY date ASC, id ASC", (month,)
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Журнал"
    headers = ["тип", "категория", "сумма, руб", "дата", "комментарий", "баланс"]
    ws.append(headers)
    hfill = PatternFill("solid", fgColor="2563EB")
    hfont = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center")

    bal = 0
    for o in ops:
        sign = 1 if o["type"] == "income" else -1
        bal += sign * o["amount"]
        tname = "Доход" if o["type"] == "income" else "Расход"
        ws.append([tname, o["category"], o["amount"], o["date"], o["comment"] or "", round(bal, 2)])

    widths = {"A": 12, "B": 22, "C": 14, "D": 14, "E": 30, "F": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = app.response_class(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["Content-Disposition"] = f"attachment; filename=journal_{month}.xlsx"
    return resp


@app.route("/delete/<int:oid>")
def delete(oid):
    conn = get_db()
    conn.execute("DELETE FROM operations WHERE id=?", (oid,))
    conn.commit()
    conn.close()
    return redirect(request.referrer or url_for("index"))


if __name__ == "__main__":
    init_db()
    app.run(host="127.0.0.1", port=5000, debug=True)